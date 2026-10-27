
import random
import re
from io import BytesIO

import pandas as pd
import streamlit as st
from docx import Document


# ==================================================
# PAGE CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="ExamPrep Randomizer",
    page_icon="📚",
    layout="centered",
)

st.title("📚 ExamPrep Randomizer")
st.write("Generate random university-style practice papers.")
st.divider()


# ==================================================
# EXAM PATTERN
# ==================================================

SECTIONS = [
    {"name": "Section A", "count": 5, "marks": 2},
    {"name": "Section B", "count": 4, "marks": 5},
    {"name": "Section C", "count": 2, "marks": 10},
]

TOTAL_QUESTIONS = sum(s["count"] for s in SECTIONS)
TOTAL_MARKS = sum(s["count"] * s["marks"] for s in SECTIONS)


# ==================================================
# QUESTION BANK HELPERS
# ==================================================

QUESTION_NAMES = {
    "question", "questions", "questiontext",
    "questionstatement", "q", "problem", "problems",
}

ANSWER_NAMES = {
    "answer", "answers", "solution", "solutions",
    "correctanswer", "answertext", "ans",
}


def normalize(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def make_question(question, answer=""):
    """Create a standardized question record."""
    if pd.isna(question):
        return None

    question = str(question).strip()

    if not question:
        return None

    if answer is None or pd.isna(answer):
        answer = ""
    else:
        answer = str(answer).strip()

    return {
        "question": question,
        "answer": answer,
    }


def deduplicate(questions):
    """Remove exact duplicate question-answer pairs."""
    result = []
    seen = set()

    for item in questions:
        key = (
            item["question"].strip().casefold(),
            item["answer"].strip().casefold(),
        )

        if key not in seen:
            seen.add(key)
            result.append(item)

    return result


# ==================================================
# EXCEL READER
# ==================================================

def read_excel_questions(file):
    """Read questions from .xlsx or .xls worksheets."""
    sheets = pd.read_excel(
        file,
        sheet_name=None,
        header=None,
    )

    questions = []

    for _, df in sheets.items():
        if df.empty:
            continue

        headings = [
            normalize(value)
            for value in df.iloc[0].tolist()
        ]

        question_col = next(
            (
                i for i, name in enumerate(headings)
                if name in QUESTION_NAMES
            ),
            None,
        )

        answer_col = next(
            (
                i for i, name in enumerate(headings)
                if name in ANSWER_NAMES
            ),
            None,
        )

        if question_col is not None:
            data = df.iloc[1:]
        else:
            # Without headings, assume:
            # Column A = question, Column B = answer.
            question_col = 0
            answer_col = 1 if df.shape[1] > 1 else None
            data = df

        for _, row in data.iterrows():
            question = row.iloc[question_col]

            answer = (
                row.iloc[answer_col]
                if answer_col is not None
                else ""
            )

            item = make_question(question, answer)

            if item:
                questions.append(item)

    return deduplicate(questions)


# ==================================================
# WORD READER
# ==================================================

QUESTION_START = re.compile(
    r"^\s*(?:(?:question|q)\s*)?\d+\s*[\).:\-]\s*(.+)$",
    re.IGNORECASE,
)

ANSWER_START = re.compile(
    r"^\s*(?:answer|ans|solution)\s*[:\-]\s*(.*)$",
    re.IGNORECASE,
)


def read_word_table(table):
    """Read a table with Question and optional Answer columns."""
    rows = [
        [cell.text.strip() for cell in row.cells]
        for row in table.rows
    ]

    if not rows:
        return []

    headings = [normalize(cell) for cell in rows[0]]

    question_col = next(
        (
            i for i, name in enumerate(headings)
            if name in QUESTION_NAMES
        ),
        None,
    )

    answer_col = next(
        (
            i for i, name in enumerate(headings)
            if name in ANSWER_NAMES
        ),
        None,
    )

    if question_col is not None:
        rows = rows[1:]
    else:
        question_col = 0
        answer_col = 1 if len(rows[0]) > 1 else None

    result = []

    for row in rows:
        if question_col >= len(row):
            continue

        answer = (
            row[answer_col]
            if answer_col is not None and answer_col < len(row)
            else ""
        )

        item = make_question(row[question_col], answer)

        if item:
            result.append(item)

    return result


def read_word_paragraphs(document):
    """Read numbered questions and their labeled answers."""
    lines = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip()
    ]

    result = []
    current_question = None
    current_answer = []
    answer_started = False

    def save_current():
        if current_question:
            item = make_question(
                current_question,
                "\n".join(current_answer),
            )
            if item:
                result.append(item)

    numbered_format = any(
        QUESTION_START.match(line) for line in lines
    )

    for line in lines:
        q_match = QUESTION_START.match(line)
        a_match = ANSWER_START.match(line)

        if q_match:
            save_current()
            current_question = q_match.group(1).strip()
            current_answer = []
            answer_started = False

        elif a_match and current_question:
            answer_text = a_match.group(1).strip()

            if answer_text:
                current_answer.append(answer_text)

            answer_started = True

        elif current_question and answer_started:
            current_answer.append(line)

        elif current_question and numbered_format:
            current_question += "\n" + line

        else:
            # Fallback: each unnumbered paragraph is a question.
            save_current()
            current_question = line
            current_answer = []
            answer_started = False

    save_current()

    return result


def read_word_questions(file):
    """Read .docx questions from tables and paragraphs."""
    document = Document(file)
    questions = []

    for table in document.tables:
        questions.extend(read_word_table(table))

    questions.extend(read_word_paragraphs(document))

    return deduplicate(questions)


# ==================================================
# CREATE WORD QUESTION PAPER
# ==================================================

def create_word_test(question_bank, selected_sections):
    """Export the three-section paper without answers."""
    document = Document()

    document.add_heading(
        "EXAMPrep Randomizer".upper(),
        level=0,
    )

    document.add_paragraph("Random Practice Question Paper")
    document.add_paragraph(
        f"Maximum Marks: {TOTAL_MARKS}"
    )
    document.add_paragraph(
        f"Total Questions: {TOTAL_QUESTIONS}"
    )
    document.add_paragraph("Name: __________________________")
    document.add_paragraph("Date: ___________________________")
    document.add_paragraph("")

    for section in SECTIONS:
        name = section["name"]
        marks = section["marks"]
        selected = selected_sections[name]

        document.add_heading(
            f"{name} — {marks} Marks Each",
            level=1,
        )

        document.add_paragraph(
            f"Attempt all {len(selected)} questions. "
            f"Section total: {len(selected) * marks} marks."
        )

        for number, index in enumerate(selected, start=1):
            question = question_bank[index]["question"]

            paragraph = document.add_paragraph()
            paragraph.add_run(
                f"{number}. {question} "
            )
            paragraph.add_run(f"[{marks} marks]").bold = True

            document.add_paragraph("")

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()


# ==================================================
# FILE UPLOAD
# ==================================================

uploaded_file = st.file_uploader(
    "Upload your question bank",
    type=["xlsx", "xls", "docx"],
    help=(
        "Excel: questions in rows. "
        "Word: numbered questions or question tables."
    ),
)

if uploaded_file is not None:
    file_signature = (
        uploaded_file.name,
        uploaded_file.size,
    )

    if st.session_state.get("file_signature") != file_signature:
        try:
            if uploaded_file.name.lower().endswith(
                (".xlsx", ".xls")
            ):
                question_bank = read_excel_questions(
                    uploaded_file
                )
            else:
                question_bank = read_word_questions(
                    uploaded_file
                )

            st.session_state["question_bank"] = question_bank
            st.session_state["file_signature"] = file_signature
            st.session_state["selected_sections"] = {}
            st.session_state["set_id"] = 0
            st.session_state["revealed"] = set()

        except Exception as error:
            st.error(
                "Could not read this file. Check its format "
                "and make sure it is not corrupted."
            )
            st.caption(f"Technical details: {error}")
            st.stop()

    question_bank = st.session_state.get(
        "question_bank", []
    )

    st.success(
        f"Successfully loaded {len(question_bank)} questions."
    )

    if not question_bank:
        st.warning(
            "No questions were detected. Check your file's "
            "column headings or question formatting."
        )
        st.stop()

    # --------------------------------------------------
    # SHOW EXAM PATTERN
    # --------------------------------------------------

    st.subheader("Your Exam Pattern")

    col_a, col_b, col_c = st.columns(3)

    col_a.metric("Section A", "5 × 2", "10 marks")
    col_b.metric("Section B", "4 × 5", "20 marks")
    col_c.metric("Section C", "2 × 10", "20 marks")

    st.caption(
        f"Total: {TOTAL_QUESTIONS} questions · "
        f"{TOTAL_MARKS} marks"
    )

    if len(question_bank) < TOTAL_QUESTIONS:
        st.error(
            f"This exam needs at least {TOTAL_QUESTIONS} "
            f"unique questions, but your file contains only "
            f"{len(question_bank)}. Upload a larger question bank."
        )
        st.stop()

    # --------------------------------------------------
    # GENERATE / SHUFFLE
    # --------------------------------------------------

    if st.button(
        "🎲 Generate Test / Shuffle",
        type="primary",
        use_container_width=True,
    ):
        previous_sections = st.session_state.get(
            "selected_sections", {}
        )

        previous_indices = {
            index
            for indices in previous_sections.values()
            for index in indices
        }

        remaining = [
            i for i in range(len(question_bank))
            if i not in previous_indices
        ]

        # Prefer an entirely fresh paper whenever possible.
        if len(remaining) >= TOTAL_QUESTIONS:
            pool = remaining
        else:
            pool = list(range(len(question_bank)))

        picked = random.sample(pool, TOTAL_QUESTIONS)

        new_sections = {}
        start = 0

        for section in SECTIONS:
            end = start + section["count"]

            new_sections[section["name"]] = picked[start:end]

            start = end

        st.session_state["selected_sections"] = new_sections
        st.session_state["set_id"] = (
            st.session_state.get("set_id", 0) + 1
        )
        st.session_state["revealed"] = set()

    # --------------------------------------------------
    # DISPLAY GENERATED PAPER
    # --------------------------------------------------

    selected_sections = st.session_state.get(
        "selected_sections", {}
    )

    if selected_sections:
        st.divider()
        st.header("📝 Your Practice Question Paper")

        st.write(
            f"**Maximum Marks:** {TOTAL_MARKS} | "
            f"**Total Questions:** {TOTAL_QUESTIONS}"
        )

        set_id = st.session_state.get("set_id", 0)

        for section in SECTIONS:
            name = section["name"]
            marks = section["marks"]
            indices = selected_sections[name]

            st.subheader(
                f"{name} — {marks} Marks Each"
            )

            st.caption(
                f"Attempt all {len(indices)} questions. "
                f"Section total: {len(indices) * marks} marks."
            )

            for number, index in enumerate(indices, start=1):
                item = question_bank[index]

                st.markdown(
                    f"**Q{number}.** {item['question']} "
                    f"**[{marks} marks]**"
                )

                reveal_key = (set_id, index)

                if item["answer"]:
                    if st.button(
                        "👁️ Reveal Answer",
                        key=f"reveal_{set_id}_{index}",
                    ):
                        st.session_state["revealed"].add(
                            reveal_key
                        )

                    if reveal_key in st.session_state["revealed"]:
                        st.info(
                            f"**Answer:** {item['answer']}"
                        )
                else:
                    st.caption(
                        "No answer was provided for this question."
                    )

                st.divider()

        # --------------------------------------------------
        # DOWNLOAD PAPER
        # --------------------------------------------------

        word_bytes = create_word_test(
            question_bank,
            selected_sections,
        )

        st.download_button(
            "⬇️ Download Question Paper (Word)",
            data=word_bytes,
            file_name="exam_practice_paper.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True,
        )

else:
    st.info(
        "Upload an Excel or Word question bank to get started."
    )

    with st.expander("How to format your question bank"):
        st.markdown(
            """
            **Excel**

            | Question | Answer |
            |---|---|
            | Define force. | A push or pull. |
            | Define velocity. | Rate of change of displacement. |

            You can omit the Answer column.

            **Word**

            1. Define force.

            Answer: A push or pull.

            2. Define velocity.

            Answer: Rate of change of displacement.

            Alternatively, use a Word table with `Question`
            and `Answer` headings.
            """
        )
